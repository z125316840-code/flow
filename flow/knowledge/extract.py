# Copyright (c) 2026, Frappe Technologies and contributors
# License: MIT. See LICENSE

"""Turn a Flow Knowledge Source into plain text, ready for chunking."""

from __future__ import annotations

import os
import re
import threading
from dataclasses import dataclass
from io import BytesIO

import frappe
from frappe import _

TEXT_EXTENSIONS = {"txt", "text", "md", "markdown", "csv", "tsv", "log", "rst", "json", "yaml", "yml"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "bmp", "tiff", "tif", "gif"}
FILE_EXTENSIONS = {"pdf", "xlsx", "docx", "html", "htm"} | TEXT_EXTENSIONS | IMAGE_EXTENSIONS

OCR_DPI = 200

PDFIUM_LOCK = threading.Lock()

CHILD_FIELDTYPES = {"Table", "Table MultiSelect"}
HTML_FIELDTYPES = {"Text Editor"}
MARKDOWN_FIELDTYPES = {"Markdown Editor"}

URL_TIMEOUT = 20
MAX_FETCH_BYTES = 10 * 1024 * 1024


@dataclass
class ExtractedDoc:
	text: str
	reference_doctype: str | None = None
	reference_name: str | None = None


def extract(source) -> list[ExtractedDoc]:
	handlers = {
		"Text": _extract_text_source,
		"File": _extract_file_source,
		"URL": _extract_url_source,
		"DocType": _extract_doctype_source,
	}
	handler = handlers.get(source.source_type)
	if handler is None:
		frappe.throw(_("Unsupported source type: {0}").format(source.source_type))
	return handler(source)


def _extract_text_source(source) -> list[ExtractedDoc]:
	return _single((source.content or "").strip())


def _extract_file_source(source) -> list[ExtractedDoc]:
	file_doc = frappe.get_doc("File", {"file_url": source.file})
	return _single(extract_file(file_doc))


def extract_file(file_doc) -> str:
	"""Extract plain text from a File doc by extension. Usable by any caller that holds a File doc."""
	extension = os.path.splitext(file_doc.file_name or file_doc.file_url or "")[1].lower().lstrip(".")
	return _extract_by_extension(file_doc.get_content(), extension).strip()


def _extract_url_source(source) -> list[ExtractedDoc]:
	import requests

	url = (source.url or "").strip()
	_validate_public_url(url)
	response = requests.get(
		url, timeout=URL_TIMEOUT, stream=True, headers={"User-Agent": "Flow-Knowledge/1.0"}
	)
	response.raise_for_status()

	raw = _read_capped(response)
	content_type = response.headers.get("Content-Type", "").lower()
	if "application/pdf" in content_type or url.split("?")[0].lower().endswith(".pdf"):
		text = _extract_pdf(raw)
	else:
		text = _extract_html(_as_text(raw))
	return _single(text.strip())


def resolve_content_fields(meta, raw) -> list[str]:
	"""Parse a comma-separated content_fields string and validate each name
	against the doctype's meta (blocks injection via crafted field names)."""
	fields = [f.strip() for f in (raw or "").split(",") if f.strip()]
	if not fields:
		frappe.throw(_("No content fields configured."), title=_("Invalid Source"))

	allowed = {df.fieldname for df in meta.fields} | {"name", "creation", "modified", "owner"}
	unknown = [f for f in fields if f not in allowed]
	if unknown:
		frappe.throw(
			_("Unknown fields for {0}: {1}").format(meta.name, ", ".join(unknown)),
			title=_("Invalid Content Fields"),
		)

	child = [f for f in fields if (df := meta.get_field(f)) and df.fieldtype in CHILD_FIELDTYPES]
	if child:
		frappe.throw(
			_("Child table fields are not supported as content fields: {0}").format(", ".join(child)),
			title=_("Invalid Content Fields"),
		)
	return fields


def _extract_doctype_source(source) -> list[ExtractedDoc]:
	doctype = source.reference_doctype
	meta = frappe.get_meta(doctype)
	fields = resolve_content_fields(meta, source.content_fields)

	rows = frappe.get_all(
		doctype, filters=_parse_filters(source.filters), fields=["name", *fields], limit_page_length=0
	)
	docs = []
	for row in rows:
		text = _row_to_text(row, fields, meta)
		if text:
			docs.append(ExtractedDoc(text=text, reference_doctype=doctype, reference_name=row["name"]))
	return docs


# --- format extractors -------------------------------------------------------


def _extract_by_extension(content, extension: str) -> str:
	if extension == "pdf":
		return _extract_pdf(_as_bytes(content))
	if extension == "xlsx":
		return _extract_xlsx(_as_bytes(content))
	if extension == "docx":
		return _extract_docx(_as_bytes(content))
	if extension in ("html", "htm"):
		return _extract_html(_as_text(content))
	if extension in IMAGE_EXTENSIONS:
		return _ocr_image(_as_bytes(content))
	if extension in TEXT_EXTENSIONS:
		return _as_text(content)
	frappe.throw(_("Unsupported file format: .{0}").format(extension or "?"), title=_("Unsupported Format"))


def _extract_pdf(data: bytes) -> str:
	import pdfplumber
	from pdfminer.pdfdocument import PDFPasswordIncorrect

	try:
		with PDFIUM_LOCK, pdfplumber.open(BytesIO(data)) as pdf:
			pages = [_extract_pdf_page(page) for page in pdf.pages]
	except PDFPasswordIncorrect:
		frappe.throw(_("PDF is password protected and cannot be read."), title=_("Cannot Read PDF"))
	return "\n\n".join(page for page in pages if page)


def _extract_pdf_page(page) -> str:
	parts = []

	tables = page.find_tables()
	for table in tables:
		markdown = _table_to_markdown(table.extract())
		if markdown:
			parts.append(markdown)

	body = page
	for table in tables:
		body = body.outside_bbox(table.bbox)
	prose = _clean_pdf_text(body.extract_text() or "")
	if prose:
		parts.append(prose)

	# No text layer → raw scan; OCR the full page. Otherwise OCR every image region
	# individually (logos, stamps, embedded scans). Duplicated text on searchable PDFs
	# is acceptable; missing text is not.
	if page.images and not parts:
		parts.append(_ocr_image(_render_page_png(page)))
	elif page.images:
		for image in page.images:
			text = _ocr_region(page, image)
			if text:
				parts.append(text)

	return "\n\n".join(part for part in parts if part)


def _table_to_markdown(table) -> str:
	rows = [row for row in (table or []) if row]
	if not rows:
		return ""
	lines = []
	for i, row in enumerate(rows):
		cells = [str(cell or "").replace("\n", " ").replace("|", "\\|").strip() for cell in row]
		lines.append("| " + " | ".join(cells) + " |")
		if i == 0:
			lines.append("| " + " | ".join("---" for _ in cells) + " |")
	return "\n".join(lines)


def _ocr_region(page, image) -> str:
	bbox = (
		max(image["x0"], page.bbox[0]),
		max(image["top"], page.bbox[1]),
		min(image["x1"], page.bbox[2]),
		min(image["bottom"], page.bbox[3]),
	)
	if bbox[2] <= bbox[0] or bbox[3] <= bbox[1]:
		return ""
	return _ocr_image(_render_page_png(page.crop(bbox)))


def _render_page_png(page) -> bytes:
	buffer = BytesIO()
	page.to_image(resolution=OCR_DPI).original.save(buffer, format="PNG")
	return buffer.getvalue()


_ocr_engine_instance = None


def _ocr_engine():
	"""Lazily initialised singleton — RapidOCR model load is expensive."""
	global _ocr_engine_instance
	if _ocr_engine_instance is None:
		from rapidocr import RapidOCR

		_ocr_engine_instance = RapidOCR()
	return _ocr_engine_instance


def _ocr_image(data: bytes) -> str:
	result = _ocr_engine()(data)
	if not result.txts:
		return ""
	return "\n".join(result.txts)


def _extract_xlsx(data: bytes) -> str:
	import openpyxl

	workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
	try:
		sheets = []
		for worksheet in workbook.worksheets:
			rows = []
			for row in worksheet.iter_rows(values_only=True):
				cells = [str(cell) for cell in row if cell is not None]
				if cells:
					rows.append("\t".join(cells))
			if rows:
				sheets.append("# {0}\n{1}".format(worksheet.title, "\n".join(rows)))
		return "\n\n".join(sheets)
	finally:
		workbook.close()


def _extract_docx(data: bytes) -> str:
	from docx import Document

	document = Document(BytesIO(data))
	paragraphs = [p.text.strip() for p in document.paragraphs]
	return "\n\n".join(p for p in paragraphs if p)


def _extract_html(markup: str) -> str:
	from bs4 import BeautifulSoup

	soup = BeautifulSoup(markup, "lxml")
	for tag in soup(["script", "style", "noscript"]):
		tag.decompose()
	return soup.get_text(separator=" ", strip=True)


# --- helpers -----------------------------------------------------------------


def _single(text: str) -> list[ExtractedDoc]:
	return [ExtractedDoc(text=text)] if text else []


def _as_bytes(content) -> bytes:
	return content if isinstance(content, bytes) else content.encode("utf-8", "replace")


def _as_text(content) -> str:
	if isinstance(content, str):
		return content
	import chardet

	encoding = (chardet.detect(content) or {}).get("encoding") or "utf-8"
	return content.decode(encoding, errors="replace")


def _clean_pdf_text(text: str) -> str:
	"""Join single newlines (line-break artifacts) while keeping paragraph breaks."""
	text = re.sub(r"(?<!\n)[ \t]*\n[ \t]*(?!\n)", " ", text)
	text = re.sub(r"[ \t]{2,}", " ", text)
	return text.strip()


def _parse_filters(raw) -> dict:
	if not raw:
		return {}
	import json

	try:
		parsed = json.loads(raw)
	except (TypeError, ValueError):
		frappe.throw(_("Filters must be valid JSON."), title=_("Invalid Filters"))
	if not isinstance(parsed, dict):
		frappe.throw(_("Filters must be a JSON object."), title=_("Invalid Filters"))
	return parsed


def _row_to_text(row: dict, fields: list[str], meta) -> str:
	from frappe.utils import cstr

	lines = []
	for fieldname in fields:
		value = row.get(fieldname)
		if value in (None, ""):
			continue
		df = meta.get_field(fieldname)
		text = _render_rich_text(df, cstr(value))
		if not text:
			continue
		label = df.label if df else fieldname
		lines.append(f"{label}: {text}")
	return "\n".join(lines)


def _render_rich_text(df, value: str) -> str:
	"""Strip HTML/Markdown markup to plain text; pass other fieldtypes through unchanged."""
	if not df:
		return value
	if df.fieldtype in HTML_FIELDTYPES:
		return _extract_html(value)
	if df.fieldtype in MARKDOWN_FIELDTYPES or (df.fieldtype == "Code" and df.options == "Markdown"):
		from frappe.utils import md_to_html

		return _extract_html(md_to_html(value))
	return value


def _validate_public_url(url: str) -> None:
	"""Block non-http(s) schemes and hosts that resolve to internal addresses (SSRF)."""
	import ipaddress
	import socket
	from urllib.parse import urlparse

	parsed = urlparse(url)
	if parsed.scheme not in ("http", "https"):
		frappe.throw(_("URL must use http or https."), title=_("Invalid URL"))
	if not parsed.hostname:
		frappe.throw(_("URL has no host."), title=_("Invalid URL"))

	try:
		infos = socket.getaddrinfo(parsed.hostname, None)
	except socket.gaierror:
		frappe.throw(_("Could not resolve URL host."), title=_("Invalid URL"))

	for info in infos:
		ip = ipaddress.ip_address(info[4][0])
		if (
			ip.is_private
			or ip.is_loopback
			or ip.is_link_local
			or ip.is_reserved
			or ip.is_multicast
			or ip.is_unspecified
		):
			frappe.throw(_("URL resolves to a non-public address."), title=_("Blocked URL"))


def _read_capped(response) -> bytes:
	chunks = []
	total = 0
	for chunk in response.iter_content(8192):
		total += len(chunk)
		if total > MAX_FETCH_BYTES:
			frappe.throw(_("Remote file exceeds the size limit."), title=_("File Too Large"))
		chunks.append(chunk)
	return b"".join(chunks)
